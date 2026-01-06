#!/usr/bin/env python3
"""
Excel writer for LinkedIn job search results.
Usage:
  python excel_writer.py <excel_file> --batch <base64_json_array>
  python excel_writer.py <excel_file> --batch-file <file_with_base64>
  python excel_writer.py <excel_file> --log <base64_json_log>
  python excel_writer.py <excel_file> --raw <json_file>
"""

import sys
import os
import json
import base64
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def write_batch(excel_file: str, jobs_data: list):
    """
    Write multiple job log entries at once to a new Excel file.
    """
    wb = Workbook()
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    ws = wb.create_sheet('Job Log', 0)

    # Define log columns
    log_cols = [
        'TIMESTAMP', 'SEARCH TITLE', 'JOB TITLE', 'COMPANY',
        'LINK', 'DECISION', 'REASON', 'SCORE'
    ]

    # Write header row
    header_fill = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, col_name in enumerate(log_cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Set column widths
    widths = [20, 25, 40, 30, 12, 15, 50, 8]
    for col_idx, width in enumerate(widths, 1):
        col_letter = chr(64 + col_idx) if col_idx <= 26 else f"A{chr(64 + col_idx - 26)}"
        ws.column_dimensions[col_letter].width = width

    # Freeze header row
    ws.freeze_panes = 'A2'

    # Decision colors
    decision_colors = {
        'MATCHED': '90EE90',       # Light green
        'REJECTED_TITLE': 'FFB6C1', # Light pink
        'REJECTED_DESC': 'FFB6C1',  # Light pink
        'REJECTED_AI': 'FFDAB9',    # Peach
        'REVIEW': 'FFD700',         # Gold
        'SKIPPED': 'D3D3D3',        # Light gray
    }

    # Write all jobs
    for row_idx, log_data in enumerate(jobs_data, 2):
        decision = log_data.get('decision', 'UNKNOWN')
        decision_color = decision_colors.get(decision, 'FFFFFF')

        row_values = [
            log_data.get('timestamp', datetime.now().isoformat()),
            log_data.get('search_title', ''),
            log_data.get('job_title', ''),
            log_data.get('company', ''),
            log_data.get('job_link', ''),
            decision,
            log_data.get('reason', ''),
            log_data.get('score', '')
        ]

        for col_idx, value in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border

            # Make link a clickable hyperlink (column 5)
            if col_idx == 5 and value:
                cell.value = 'Open'
                cell.hyperlink = value
                cell.font = Font(color='0563C1', underline='single')
            else:
                cell.value = value

            # Color the decision column (column 6)
            if col_idx == 6:
                cell.fill = PatternFill(start_color=decision_color, end_color=decision_color, fill_type='solid')
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')

    # Save
    wb.save(excel_file)
    print(f"Saved {len(jobs_data)} jobs to {excel_file}")


def write_log(excel_file: str, log_data: dict):
    """
    Write a single job decision log entry to the Job Log sheet.
    """
    # Check if file exists
    if os.path.exists(excel_file):
        wb = load_workbook(excel_file)
    else:
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

    # Get or create Job Log sheet
    if 'Job Log' in wb.sheetnames:
        ws = wb['Job Log']
    else:
        ws = wb.create_sheet('Job Log', 0)

        # Define log columns
        log_cols = [
            'TIMESTAMP', 'SEARCH TITLE', 'JOB TITLE', 'COMPANY',
            'LINK', 'DECISION', 'REASON', 'SCORE'
        ]

        # Write header row
        header_fill = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for col_idx, col_name in enumerate(log_cols, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        # Set column widths
        widths = [20, 25, 40, 30, 12, 15, 50, 8]
        for col_idx, width in enumerate(widths, 1):
            col_letter = chr(64 + col_idx) if col_idx <= 26 else f"A{chr(64 + col_idx - 26)}"
            ws.column_dimensions[col_letter].width = width

        # Freeze header row
        ws.freeze_panes = 'A2'

    # Find next empty row
    next_row = ws.max_row + 1

    # Decision colors
    decision = log_data.get('decision', 'UNKNOWN')
    decision_colors = {
        'MATCHED': '90EE90',
        'REJECTED_TITLE': 'FFB6C1',
        'REJECTED_DESC': 'FFB6C1',
        'REJECTED_AI': 'FFDAB9',
        'REVIEW': 'FFD700',
        'SKIPPED': 'D3D3D3',
    }
    decision_color = decision_colors.get(decision, 'FFFFFF')

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    row_values = [
        log_data.get('timestamp', datetime.now().isoformat()),
        log_data.get('search_title', ''),
        log_data.get('job_title', ''),
        log_data.get('company', ''),
        log_data.get('job_link', ''),
        decision,
        log_data.get('reason', ''),
        log_data.get('score', '')
    ]

    for col_idx, value in enumerate(row_values, 1):
        cell = ws.cell(row=next_row, column=col_idx)
        cell.border = thin_border

        if col_idx == 5 and value:
            cell.value = 'Open'
            cell.hyperlink = value
            cell.font = Font(color='0563C1', underline='single')
        else:
            cell.value = value

        if col_idx == 6:
            cell.fill = PatternFill(start_color=decision_color, end_color=decision_color, fill_type='solid')
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

    wb.save(excel_file)
    print(f"Saved log entry to {excel_file}")


def write_raw(excel_file: str, json_file: str):
    """
    Export raw scraped job data to Excel.
    """
    try:
        with open(json_file, 'r', encoding='utf-8') as f:
            jobs = json.load(f)
    except Exception as e:
        print(f"Error loading JSON file {json_file}: {e}")
        return

    if not jobs:
        print(f"No jobs in {json_file}")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Raw Jobs"

    raw_cols = ['TITLE', 'COMPANY', 'LOCATION', 'LINK', 'POSTED DATE', 'DESCRIPTION']

    header_fill = PatternFill(start_color='6B5B95', end_color='6B5B95', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, col_name in enumerate(raw_cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    widths = [40, 30, 25, 12, 15, 80]
    for col_idx, width in enumerate(widths, 1):
        col_letter = chr(64 + col_idx) if col_idx <= 26 else f"A{chr(64 + col_idx - 26)}"
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = 'A2'

    for row_idx, job in enumerate(jobs, 2):
        row_values = [
            job.get('title', ''),
            job.get('company', ''),
            job.get('location', ''),
            job.get('job_link', ''),
            job.get('posted_date', ''),
            job.get('description', '')[:5000] if job.get('description') else ''
        ]

        for col_idx, value in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border

            if col_idx == 4 and value:
                cell.value = 'Open'
                cell.hyperlink = value
                cell.font = Font(color='0563C1', underline='single')
            else:
                cell.value = value

    wb.save(excel_file)
    print(f"Exported {len(jobs)} raw jobs to {excel_file}")


def merge_excels(output_file: str, input_pattern: str):
    """
    Merge multiple Excel files matching a pattern into one combined file.
    Also creates a Summary sheet with counts per search title.
    """
    import glob

    # Find all matching Excel files
    input_files = sorted(glob.glob(input_pattern))

    if not input_files:
        print(f"No files found matching pattern: {input_pattern}")
        return

    print(f"Found {len(input_files)} files to merge")

    # Create output workbook
    wb = Workbook()
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    ws = wb.create_sheet('All Jobs', 0)

    # Define columns
    log_cols = [
        'TIMESTAMP', 'SEARCH TITLE', 'JOB TITLE', 'COMPANY',
        'LINK', 'DECISION', 'REASON', 'SCORE'
    ]

    # Write header row
    header_fill = PatternFill(start_color='1565C0', end_color='1565C0', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, col_name in enumerate(log_cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Set column widths
    widths = [20, 25, 40, 30, 12, 15, 50, 8]
    for col_idx, width in enumerate(widths, 1):
        col_letter = chr(64 + col_idx) if col_idx <= 26 else f"A{chr(64 + col_idx - 26)}"
        ws.column_dimensions[col_letter].width = width

    ws.freeze_panes = 'A2'

    # Decision colors
    decision_colors = {
        'MATCHED': '90EE90',
        'REJECTED_TITLE': 'FFB6C1',
        'REJECTED_DESC': 'FFB6C1',
        'REJECTED_AI': 'FFDAB9',
        'REVIEW': 'FFD700',
        'SKIPPED': 'D3D3D3',
    }

    # Track stats per search title
    stats = {}
    current_row = 2
    total_jobs = 0

    # Read each input file and append rows
    for input_file in input_files:
        try:
            src_wb = load_workbook(input_file)
            src_ws = src_wb.active

            # Skip header row, read all data rows
            for src_row in range(2, src_ws.max_row + 1):
                row_data = []
                for col in range(1, 9):  # 8 columns
                    cell_value = src_ws.cell(row=src_row, column=col).value
                    # Get hyperlink if it's a link column
                    if col == 5:
                        cell_obj = src_ws.cell(row=src_row, column=col)
                        if cell_obj.hyperlink:
                            row_data.append(cell_obj.hyperlink.target)
                        else:
                            row_data.append(cell_value or '')
                    else:
                        row_data.append(cell_value or '')

                if not any(row_data):  # Skip empty rows
                    continue

                # Track stats
                search_title = row_data[1]
                decision = row_data[5]
                if search_title not in stats:
                    stats[search_title] = {'total': 0, 'matched': 0, 'rejected': 0, 'skipped': 0}
                stats[search_title]['total'] += 1
                if decision == 'MATCHED':
                    stats[search_title]['matched'] += 1
                elif decision == 'SKIPPED':
                    stats[search_title]['skipped'] += 1
                else:
                    stats[search_title]['rejected'] += 1

                # Write row to merged sheet
                decision_color = decision_colors.get(decision, 'FFFFFF')

                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.border = thin_border

                    if col_idx == 5 and value:  # Link column
                        cell.value = 'Open'
                        cell.hyperlink = value
                        cell.font = Font(color='0563C1', underline='single')
                    else:
                        cell.value = value

                    if col_idx == 6:  # Decision column
                        cell.fill = PatternFill(start_color=decision_color, end_color=decision_color, fill_type='solid')
                        cell.font = Font(bold=True)
                        cell.alignment = Alignment(horizontal='center')

                current_row += 1
                total_jobs += 1

            src_wb.close()
            print(f"  Merged: {os.path.basename(input_file)}")

        except Exception as e:
            print(f"  Error reading {input_file}: {e}")

    # Create Summary sheet
    summary_ws = wb.create_sheet('Summary', 0)

    summary_cols = ['SEARCH TITLE', 'TOTAL JOBS', 'MATCHED', 'REJECTED', 'SKIPPED', 'MATCH RATE']
    summary_fill = PatternFill(start_color='2E7D32', end_color='2E7D32', fill_type='solid')

    for col_idx, col_name in enumerate(summary_cols, 1):
        cell = summary_ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = summary_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # Set summary column widths
    summary_widths = [30, 12, 10, 10, 10, 12]
    for col_idx, width in enumerate(summary_widths, 1):
        col_letter = chr(64 + col_idx)
        summary_ws.column_dimensions[col_letter].width = width

    summary_ws.freeze_panes = 'A2'

    # Write summary rows
    summary_row = 2
    total_matched = 0
    total_rejected = 0
    total_skipped = 0

    for search_title, counts in sorted(stats.items()):
        total = counts['total']
        matched = counts['matched']
        rejected = counts['rejected']
        skipped = counts['skipped']
        match_rate = f"{(matched/total*100):.1f}%" if total > 0 else "0%"

        total_matched += matched
        total_rejected += rejected
        total_skipped += skipped

        row_values = [search_title, total, matched, rejected, skipped, match_rate]

        for col_idx, value in enumerate(row_values, 1):
            cell = summary_ws.cell(row=summary_row, column=col_idx)
            cell.value = value
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center' if col_idx > 1 else 'left')

            # Color match rate
            if col_idx == 6:
                rate_val = matched / total * 100 if total > 0 else 0
                if rate_val >= 20:
                    cell.fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')
                elif rate_val >= 10:
                    cell.fill = PatternFill(start_color='FFD700', end_color='FFD700', fill_type='solid')

        summary_row += 1

    # Add totals row
    total_rate = f"{(total_matched/total_jobs*100):.1f}%" if total_jobs > 0 else "0%"
    totals = ['TOTAL', total_jobs, total_matched, total_rejected, total_skipped, total_rate]

    for col_idx, value in enumerate(totals, 1):
        cell = summary_ws.cell(row=summary_row, column=col_idx)
        cell.value = value
        cell.border = thin_border
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center' if col_idx > 1 else 'left')
        cell.fill = PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')

    # Save
    wb.save(output_file)
    print(f"\nMerged {total_jobs} jobs from {len(input_files)} files into {output_file}")
    print(f"Summary: {total_matched} matched, {total_rejected} rejected, {total_skipped} skipped")


def main():
    if len(sys.argv) < 3:
        print("Usage: python excel_writer.py <excel_file> --batch <base64_json_array>")
        print("       python excel_writer.py <excel_file> --batch-file <file_with_base64>")
        print("       python excel_writer.py <excel_file> --log <base64_json_log>")
        print("       python excel_writer.py <excel_file> --raw <json_file>")
        print("       python excel_writer.py <output_file> --merge <input_pattern>")
        sys.exit(1)

    excel_file = sys.argv[1]

    # Ensure output directory exists
    os.makedirs(os.path.dirname(excel_file), exist_ok=True)

    if sys.argv[2] == '--raw':
        if len(sys.argv) < 4:
            print("Error: --raw requires JSON file path")
            sys.exit(1)
        write_raw(excel_file, sys.argv[3])
        return

    if sys.argv[2] == '--log':
        if len(sys.argv) < 4:
            print("Error: --log requires base64 JSON data")
            sys.exit(1)
        try:
            json_str = base64.b64decode(sys.argv[3]).decode('utf-8')
            log_data = json.loads(json_str)
        except Exception as e:
            print(f"Error decoding log data: {e}")
            sys.exit(1)
        write_log(excel_file, log_data)
        return

    if sys.argv[2] == '--batch':
        if len(sys.argv) < 4:
            print("Error: --batch requires base64 JSON array")
            sys.exit(1)
        try:
            json_str = base64.b64decode(sys.argv[3]).decode('utf-8')
            jobs_data = json.loads(json_str)
            if not isinstance(jobs_data, list):
                print("Error: --batch data must be a JSON array")
                sys.exit(1)
        except Exception as e:
            print(f"Error decoding batch data: {e}")
            sys.exit(1)
        write_batch(excel_file, jobs_data)
        return

    if sys.argv[2] == '--batch-file':
        if len(sys.argv) < 4:
            print("Error: --batch-file requires file path")
            sys.exit(1)
        batch_file = sys.argv[3]
        try:
            with open(batch_file, 'r') as f:
                base64_data = f.read().strip()
            json_str = base64.b64decode(base64_data).decode('utf-8')
            jobs_data = json.loads(json_str)
            if not isinstance(jobs_data, list):
                print("Error: --batch-file data must be a JSON array")
                sys.exit(1)
        except Exception as e:
            print(f"Error reading/decoding batch file: {e}")
            sys.exit(1)
        write_batch(excel_file, jobs_data)
        return

    if sys.argv[2] == '--merge':
        if len(sys.argv) < 4:
            print("Error: --merge requires input file pattern")
            sys.exit(1)
        input_pattern = sys.argv[3]
        merge_excels(excel_file, input_pattern)
        return

    print(f"Unknown option: {sys.argv[2]}")
    sys.exit(1)


if __name__ == "__main__":
    main()
